"""
DOCX report generation using python-docx.

Imports shared color constants and _build_inline_segments from report_service.
Provides:
  generate_docx_report()   — redline-style full document report
  generate_docx_summary()  — analytics summary (score, tables, jurisdiction)
"""
import io
import logging

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import lxml.etree as etree
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False
    logger.warning("python-docx not installed — DOCX generation unavailable")

from .report_service import (
    _build_inline_segments,
    _get_risk_info,
    CATEGORY_ORDER,
    _CATEGORY_RISK_LEVEL,
    COLOR_RED, COLOR_GREEN, COLOR_ORANGE, COLOR_BLUE,
    COLOR_BLACK, COLOR_GREY, COLOR_DARK, COLOR_RULE,
    BG_RED, BG_GREEN, BG_ORANGE, BG_BLUE,
    COLOR_SCORE_HIGH, COLOR_SCORE_MED, COLOR_SCORE_LOW,
)


# ── Color helpers ─────────────────────────────────────────────────────────────

def _hex_to_rgb(h: str) -> tuple:
    """Convert #rrggbb → (r, g, b) int tuple."""
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _rgb(hex_color: str) -> "RGBColor":
    """Convert #rrggbb to python-docx RGBColor."""
    r, g, b = _hex_to_rgb(hex_color)
    return RGBColor(r, g, b)


def _set_cell_shading(cell, hex_color: str):
    """Apply background fill to a table cell via OOXML."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.lstrip("#"))
    tcPr.append(shd)


def _set_paragraph_shading(paragraph, hex_color: str):
    """Apply background fill to a paragraph via OOXML pPr/shd."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.lstrip("#"))
    pPr.append(shd)


def _docx_dark_header_table(doc: "Document", headers: list, col_widths: list) -> "Table":
    """Create a one-row header table with dark background and white bold text."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, (header, width) in enumerate(zip(headers, col_widths)):
        cell = t.rows[0].cells[i]
        cell.width = width
        _set_cell_shading(cell, COLOR_DARK)
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    return t


def _docx_agreement_block(doc: "Document", agreement_meta: dict):
    """Add Agreement Details section to document."""
    if not agreement_meta:
        return

    doc.add_heading("Agreement Details", level=2)

    agmt_type = agreement_meta.get("agreement_type", "")
    agmt_det  = agreement_meta.get("agreement_details") or {}
    parties   = agreement_meta.get("parties") or {}
    prop      = agreement_meta.get("property") or {}

    def kv(label, value):
        if not value:
            return
        p = doc.add_paragraph()
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(str(value))

    kv("Agreement Type", agmt_type)
    kv("Agreement Date", agmt_det.get("agreement_date"))
    place = ", ".join(v for v in [agmt_det.get("city"), agmt_det.get("state")] if v)
    kv("Location", place)

    landlord = parties.get("landlord") or {}
    tenant   = parties.get("tenant") or {}
    if any(landlord.values()) or any(tenant.values()):
        doc.add_heading("Parties", level=3)
        kv("Landlord", landlord.get("name"))
        kv("Landlord Address", landlord.get("address"))
        kv("Landlord Contact", landlord.get("contact"))
        tenant_name = tenant.get("company_name") or tenant.get("name", "")
        kv("Tenant", tenant_name)
        kv("Authorized By", tenant.get("authorized_signatory"))
        kv("Tenant Address", tenant.get("address"))
        kv("Tenant Contact", tenant.get("contact"))

    if any(prop.values()):
        doc.add_heading("Property", level=3)
        kv("Type", prop.get("type"))
        kv("Area", f"{prop['area_sqft']} sq ft" if prop.get("area_sqft") else None)
        kv("Address", prop.get("address"))

    doc.add_paragraph("─" * 60)


# ══════════════════════════════════════════════════════════════════════════════
#  DOCX REPORT — redline-style full document
# ══════════════════════════════════════════════════════════════════════════════

def generate_docx_report(
    analysis_summary: list,
    *,
    conflicts: list = None,
    jurisdiction_info: dict = None,
    full_text: str = "",
    agreement_meta: dict = None,
) -> bytes:
    """
    Generate a redline-style DOCX report.

    Returns raw DOCX bytes. Raises RuntimeError if python-docx is not installed.
    """
    if not _DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed. Cannot generate DOCX reports.")

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    doc.add_heading("Document Analysis Report", level=1)

    if agreement_meta:
        _docx_agreement_block(doc, agreement_meta)

    doc.add_heading("Document with Annotations", level=2)

    segments = _build_inline_segments(full_text, analysis_summary)

    i = 0
    while i < len(segments):
        seg   = segments[i]
        stype = seg["type"]
        text  = seg["text"].replace("<br>", "\n").replace("<b>", "").replace("</b>", "")

        if stype == "violation":
            p = doc.add_paragraph()
            _set_paragraph_shading(p, BG_RED)
            run = p.add_run(f"− {text}")
            run.font.color.rgb = _rgb(COLOR_RED)
            run.font.strike = True
            reason = seg.get("reason", "")
            if reason:
                p.add_run(f"  [{reason}]").font.color.rgb = _rgb(COLOR_GREY)
            if i + 1 < len(segments) and segments[i + 1]["type"] == "ai":
                i += 1
                p2 = doc.add_paragraph()
                _set_paragraph_shading(p2, BG_GREEN)
                run2 = p2.add_run(f"+ {segments[i]['text']}")
                run2.font.color.rgb = _rgb(COLOR_GREEN)

        elif stype == "partial":
            p = doc.add_paragraph()
            _set_paragraph_shading(p, BG_ORANGE)
            run = p.add_run(f"~ {text}")
            run.font.color.rgb = _rgb(COLOR_ORANGE)
            run.font.strike = True
            reason = seg.get("reason", "")
            if reason:
                p.add_run(f"  [{reason}]").font.color.rgb = _rgb(COLOR_GREY)
            if i + 1 < len(segments) and segments[i + 1]["type"] == "ai":
                i += 1
                p2 = doc.add_paragraph()
                _set_paragraph_shading(p2, BG_GREEN)
                run2 = p2.add_run(f"+ {segments[i]['text']}")
                run2.font.color.rgb = _rgb(COLOR_GREEN)

        elif stype == "not_found_ai":
            p = doc.add_paragraph()
            _set_paragraph_shading(p, BG_BLUE)
            run = p.add_run(f"+ {text}")
            run.font.color.rgb = _rgb(COLOR_BLUE)

        elif stype == "match":
            p = doc.add_paragraph(text)
            p.runs[0].font.color.rgb = _rgb(COLOR_BLACK) if p.runs else None

        elif stype == "heading":
            doc.add_heading(text, level=3)

        elif stype == "bullet":
            doc.add_paragraph(text, style="List Bullet")

        elif stype == "ai":
            p = doc.add_paragraph()
            _set_paragraph_shading(p, BG_GREEN)
            run = p.add_run(f"+ {text}")
            run.font.color.rgb = _rgb(COLOR_GREEN)

        elif stype == "normal":
            doc.add_paragraph(text)

        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
#  DOCX SUMMARY — analytics dashboard
# ══════════════════════════════════════════════════════════════════════════════

def generate_docx_summary(
    analysis_summary: list,
    *,
    conflicts: list = None,
    jurisdiction_info: dict = None,
    full_text: str = "",
    agreement_meta: dict = None,
) -> bytes:
    """
    Generate a DOCX analytics summary with compliance score, tables, and jurisdiction.

    Returns raw DOCX bytes. Raises RuntimeError if python-docx is not installed.
    """
    if not _DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed. Cannot generate DOCX reports.")

    from collections import defaultdict

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    jurisdiction_info = jurisdiction_info or {}

    match_count     = sum(1 for c in analysis_summary if c["result"] == "MATCH")
    violation_count = sum(1 for c in analysis_summary if c["result"] == "VIOLATION")
    partial_count   = sum(1 for c in analysis_summary if c["result"] == "PARTIALLY_SATISFIED")
    not_found_count = sum(1 for c in analysis_summary if c["result"] == "NOT_FOUND")
    total           = len(analysis_summary)
    compliance_score = round((match_count / total) * 100) if total else 0

    if compliance_score >= 80:
        score_color  = COLOR_SCORE_HIGH
        score_status = "Compliant"
    elif compliance_score >= 50:
        score_color  = COLOR_SCORE_MED
        score_status = "Needs Attention"
    else:
        score_color  = COLOR_SCORE_LOW
        score_status = "Critical"

    doc.add_heading("Document Analysis Summary", level=1)

    if agreement_meta:
        _docx_agreement_block(doc, agreement_meta)

    # Overall Compliance Score
    score_p = doc.add_paragraph()
    score_run = score_p.add_run(f"Overall Compliance Score: {compliance_score}% — {score_status}")
    score_run.bold = True
    score_run.font.size = Pt(16)
    score_run.font.color.rgb = _rgb(score_color)
    score_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Stats table
    doc.add_heading("Statistics", level=2)
    stats_tbl = doc.add_table(rows=2, cols=5)
    stats_tbl.style = "Table Grid"
    headers = ["Total", "Match", "Violation", "Partial", "Not Found"]
    values  = [total, match_count, violation_count, partial_count, not_found_count]
    for i, h in enumerate(headers):
        cell = stats_tbl.rows[0].cells[i]
        _set_cell_shading(cell, COLOR_DARK)
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    val_colors = [COLOR_BLACK, COLOR_GREEN, COLOR_RED, COLOR_ORANGE, COLOR_BLUE]
    for i, (v, vc) in enumerate(zip(values, val_colors)):
        cell = stats_tbl.rows[1].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(str(v))
        run.bold = True
        run.font.color.rgb = _rgb(vc)
        run.font.size = Pt(11)

    doc.add_paragraph()

    # Critical issues
    doc.add_heading("Critical Issues Requiring Immediate Attention", level=2)
    red_flags = [e for e in analysis_summary if e["result"] in ("VIOLATION", "NOT_FOUND")][:3]
    if not red_flags:
        doc.add_paragraph("No critical issues found.")
    else:
        rf_tbl = _docx_dark_header_table(
            doc, ["Clause", "Status", "Issue"],
            [Cm(6), Cm(4), Cm(8)],
        )
        for flag in red_flags:
            row = rf_tbl.add_row()
            row.cells[0].text = flag["clause_title"]
            status_cell = row.cells[1]
            _set_cell_shading(status_cell, BG_RED if flag["result"] == "VIOLATION" else BG_BLUE)
            sr = status_cell.paragraphs[0].add_run(flag["result"])
            sr.bold = True
            sr.font.color.rgb = _rgb(COLOR_RED if flag["result"] == "VIOLATION" else COLOR_BLUE)
            row.cells[2].text = flag.get("reason", "")

    doc.add_paragraph()

    # Risk category breakdown
    doc.add_heading("Risk Category Breakdown", level=2)
    category_stats: dict = defaultdict(lambda: {"total": 0, "compliant": 0, "issues": 0})
    for entry in analysis_summary:
        cat, _ = _get_risk_info(entry.get("clause_title", ""))
        category_stats[cat]["total"] += 1
        if entry["result"] == "MATCH":
            category_stats[cat]["compliant"] += 1
        else:
            category_stats[cat]["issues"] += 1

    cat_tbl = _docx_dark_header_table(
        doc, ["Category", "Risk Level", "Total", "Compliant", "Issues"],
        [Cm(5), Cm(3), Cm(2), Cm(3), Cm(2)],
    )
    for cat in CATEGORY_ORDER:
        if cat not in category_stats:
            continue
        risk_lv, risk_color = _CATEGORY_RISK_LEVEL[cat]
        s = category_stats[cat]
        row = cat_tbl.add_row()
        row.cells[0].text = cat
        risk_cell = row.cells[1]
        run = risk_cell.paragraphs[0].add_run(risk_lv)
        run.bold = True
        run.font.color.rgb = _rgb(risk_color)
        for ci, val in zip([2, 3, 4], [s["total"], s["compliant"], s["issues"]]):
            row.cells[ci].text = str(val)
            row.cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Jurisdiction
    doc.add_heading("Jurisdiction & Compliance Checklist", level=2)
    if jurisdiction_info:
        juris     = jurisdiction_info.get("jurisdiction", "Unknown")
        agmt_type = jurisdiction_info.get("agreement_type", "Unknown")
        laws      = ", ".join(jurisdiction_info.get("applicable_laws", []))

        def kv(label, value):
            p = doc.add_paragraph()
            r = p.add_run(f"{label}: ")
            r.bold = True
            p.add_run(str(value))

        kv("Jurisdiction", juris)
        kv("Agreement Type", agmt_type)
        if laws:
            kv("Applicable Laws", laws)

        checklist = jurisdiction_info.get("checklist", [])
        if checklist:
            doc.add_paragraph()
            jc_tbl = _docx_dark_header_table(
                doc, ["Requirement", "Status"],
                [Cm(12), Cm(3)],
            )
            for item in checklist:
                req_label = "Required" if item.get("required") else "Optional"
                req_color = COLOR_RED if item.get("required") else COLOR_GREEN
                row = jc_tbl.add_row()
                row.cells[0].text = item.get("item", "")
                rc = row.cells[1]
                rr = rc.paragraphs[0].add_run(req_label)
                rr.bold = True
                rr.font.color.rgb = _rgb(req_color)
    else:
        doc.add_paragraph("Jurisdiction information not available.")

    doc.add_paragraph()

    # Contract Timeline
    doc.add_heading("Contract Timeline", level=2)
    timeline_rows = [
        (entry["clause_title"], dt)
        for entry in analysis_summary
        for dt in entry.get("key_dates_durations", [])
        if dt
    ]
    if not timeline_rows:
        doc.add_paragraph("No key dates or durations identified.")
    else:
        tl_tbl = _docx_dark_header_table(
            doc, ["Clause", "Timeline Item"],
            [Cm(7), Cm(9)],
        )
        for ct, ti in timeline_rows:
            row = tl_tbl.add_row()
            row.cells[0].text = ct
            row.cells[1].text = ti

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
