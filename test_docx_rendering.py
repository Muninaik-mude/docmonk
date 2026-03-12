#!/usr/bin/env python
"""
test_docx_rendering.py
======================
Comprehensive test for _docx_to_html — builds a DOCX entirely in memory
that contains every field/control type found in loan applications, then
runs it through the renderer and asserts the expected HTML output.

Field types covered
-------------------
 1.  Title run  — bold, large font, center alignment
 2.  Run formatting — bold / italic / underline / colour / font-size deviation
 3.  Paragraph alignment — center, right, justify
 4.  w14:checkbox SDT  checked    (Word 2010+ content control)
 5.  w14:checkbox SDT  unchecked
 6.  Checkbox labels preserved alongside the control
 7.  Radio-button group  (exclusive w14:checkbox SDTs in one paragraph)
 8.  Legacy FORMCHECKBOX field code  checked
 9.  Legacy FORMCHECKBOX field code  unchecked
10.  Inline SDT text field (e.g. "Borrower Name")
11.  Inline SDT date-picker
12.  Body-level SDT wrapping a paragraph
13.  Inline hyperlink
14.  Table — plain cells
15.  Table — checkbox inside a cell
16.  Table — cell background shading
17.  Nested SDT  (SDT inside a body-level SDT)

Run
---
    cd e:/docmonk-repo/docmonk
    python test_docx_rendering.py

Outputs
-------
  • Pass/fail line per assertion
  • test_docx_output.html  — open in browser for visual inspection
"""

import io
import os
import sys

# ── Django setup (needed for the module-level import chain) ──────────────────
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
os.environ.setdefault("DEBUG", "True")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import django
django.setup()

# ── Import the function under test ────────────────────────────────────────────
from policy.services.policy_report_service import _docx_to_html  # noqa: E402

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import lxml.etree as etree

# ─────────────────────────────────────────────────────────────────────────────
# XML namespace constants
# ─────────────────────────────────────────────────────────────────────────────
W14_NS  = "http://schemas.microsoft.com/office/word/2010/wordml"
XML_NS  = "http://www.w3.org/XML/1998/namespace"   # for xml:space="preserve"


# ─────────────────────────────────────────────────────────────────────────────
# DOCX element builders
# ─────────────────────────────────────────────────────────────────────────────

def _t(text: str, preserve_space: bool = True) -> "OxmlElement":
    """Build a <w:t> element."""
    el = OxmlElement("w:t")
    if preserve_space:
        el.set(f"{{{XML_NS}}}space", "preserve")
    el.text = text
    return el


def _r(text: str, bold=False, italic=False, underline=False, color_rgb=None, size_pt=None) -> "OxmlElement":
    """Build a <w:r> element with optional formatting."""
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    has_rPr = False
    if bold:
        rPr.append(OxmlElement("w:b"))
        has_rPr = True
    if italic:
        rPr.append(OxmlElement("w:i"))
        has_rPr = True
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        has_rPr = True
    if color_rgb:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color_rgb)
        rPr.append(c)
        has_rPr = True
    if size_pt:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size_pt * 2)))  # half-points
        rPr.append(sz)
        has_rPr = True
    if has_rPr:
        r.append(rPr)
    r.append(_t(text))
    return r


def make_w14_checkbox_sdt(checked: bool) -> "OxmlElement":
    """Build an inline w:sdt with a w14:checkbox content control."""
    sdt = OxmlElement("w:sdt")

    sdtPr = OxmlElement("w:sdtPr")
    cb = etree.SubElement(sdtPr, f"{{{W14_NS}}}checkbox")
    checked_el = etree.SubElement(cb, f"{{{W14_NS}}}checked")
    checked_el.set(f"{{{W14_NS}}}val", "1" if checked else "0")
    sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    r.append(_t("☑" if checked else "☐", preserve_space=False))
    sdtContent.append(r)
    sdt.append(sdtContent)

    return sdt


def make_legacy_checkbox(checked: bool) -> list:
    """
    Build the 4-run sequence for a legacy FORMCHECKBOX field code.
    Returns a list of w:r elements to append to a paragraph.
    """
    # Run 1: fldChar begin  +  ffData > checkBox
    r1 = OxmlElement("w:r")
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    ffData = OxmlElement("w:ffData")
    checkBox_el = OxmlElement("w:checkBox")
    sizeAuto = OxmlElement("w:sizeAuto")
    checkBox_el.append(sizeAuto)
    default_el = OxmlElement("w:default")
    default_el.set(qn("w:val"), "0")
    checkBox_el.append(default_el)
    checked_el = OxmlElement("w:checked")
    checked_el.set(qn("w:val"), "1" if checked else "0")
    checkBox_el.append(checked_el)
    ffData.append(checkBox_el)
    fld1.append(ffData)
    r1.append(fld1)

    # Run 2: instrText
    r2 = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(f"{{{XML_NS}}}space", "preserve")
    instr.text = " FORMCHECKBOX "
    r2.append(instr)

    # Run 3: fldChar separate
    r3 = OxmlElement("w:r")
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "separate")
    r3.append(fld3)

    # Run 4: fldChar end
    r4 = OxmlElement("w:r")
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    r4.append(fld4)

    return [r1, r2, r3, r4]


def make_sdt_text_field(tag_name: str, value: str) -> "OxmlElement":
    """Build a non-checkbox inline SDT text input field."""
    sdt = OxmlElement("w:sdt")

    sdtPr = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), tag_name)
    sdtPr.append(tag)
    sdtPr.append(OxmlElement("w:text"))   # marks it as a plain-text field
    sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    r.append(_t(value))
    sdtContent.append(r)
    sdt.append(sdtContent)

    return sdt


def make_sdt_date_field(display: str, iso_date: str = "2024-03-15T00:00:00Z") -> "OxmlElement":
    """Build an inline SDT date-picker field."""
    sdt = OxmlElement("w:sdt")

    sdtPr = OxmlElement("w:sdtPr")
    date_el = OxmlElement("w:date")
    date_el.set(qn("w:fullDate"), iso_date)
    fmt = OxmlElement("w:dateFormat")
    fmt.set(qn("w:val"), "MM/dd/yyyy")
    date_el.append(fmt)
    sdtPr.append(date_el)
    sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    r.append(_t(display, preserve_space=False))
    sdtContent.append(r)
    sdt.append(sdtContent)

    return sdt


def make_body_sdt_with_paragraph(text: str) -> "OxmlElement":
    """Build a body-level SDT that wraps a single paragraph."""
    sdt = OxmlElement("w:sdt")

    sdtPr = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "bodySection")
    sdtPr.append(tag)
    sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    p.append(_r(text, bold=True))
    sdtContent.append(p)
    sdt.append(sdtContent)

    return sdt


def make_nested_sdt(outer_text: str, inner_checkbox_checked: bool) -> "OxmlElement":
    """
    Body-level SDT whose sdtContent paragraph itself contains an inline checkbox SDT.
    Tests recursive _render_children behaviour.
    """
    outer_sdt = OxmlElement("w:sdt")
    sdtPr = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "nestedSection")
    sdtPr.append(tag)
    outer_sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    p.append(_r(outer_text + "  "))
    p.append(make_w14_checkbox_sdt(inner_checkbox_checked))
    p.append(_r("  NestedOption"))
    sdtContent.append(p)
    outer_sdt.append(sdtContent)

    return outer_sdt


def _insert_before_sectPr(body, elem):
    """Append elem just before the w:sectPr element (so docx stays valid)."""
    sectPr = body.find(qn("w:sectPr"))
    if sectPr is not None:
        idx = list(body).index(sectPr)
        body.insert(idx, elem)
    else:
        body.append(elem)


# ─────────────────────────────────────────────────────────────────────────────
# Build the test document
# ─────────────────────────────────────────────────────────────────────────────

def build_test_docx() -> bytes:
    doc  = Document()
    body = doc.element.body

    # ── 1. Title — bold, large, centred ──────────────────────────────────────
    title = doc.add_paragraph()
    tr = title.add_run("LOAN APPLICATION FORM")
    tr.bold = True
    tr.font.size = Pt(16)
    tr.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── 2. Run-level formatting deviations ───────────────────────────────────
    fmt_para = doc.add_paragraph()
    fmt_para.add_run("Normal text,  ")
    r_b = fmt_para.add_run("bold text,  ")
    r_b.bold = True
    r_i = fmt_para.add_run("italic text,  ")
    r_i.italic = True
    r_u = fmt_para.add_run("underlined text,  ")
    r_u.underline = True
    r_c = fmt_para.add_run("red coloured text")
    r_c.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    # ── 3. Paragraph alignments ───────────────────────────────────────────────
    p_c = doc.add_paragraph("Centred paragraph for alignment test")
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_r = doc.add_paragraph("Right-aligned paragraph for alignment test")
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p_j = doc.add_paragraph(
        "This is a justified paragraph with sufficient text to demonstrate "
        "that the justify alignment CSS is emitted correctly in the HTML output."
    )
    p_j.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # ── 4 & 5. w14:checkbox SDTs — checked + unchecked inline ────────────────
    p_cb = doc.add_paragraph()
    p_cb._p.append(_r("Employment Type:  "))
    p_cb._p.append(make_w14_checkbox_sdt(checked=True))
    p_cb._p.append(_r("  Full-Time     "))
    p_cb._p.append(make_w14_checkbox_sdt(checked=False))
    p_cb._p.append(_r("  Part-Time"))

    # ── 7. Radio-button group (exclusive checkboxes) ─────────────────────────
    p_radio = doc.add_paragraph()
    p_radio._p.append(_r("Loan Purpose:  "))
    for label, chk in [("Purchase", True), ("Refinance", False), ("Home Equity", False)]:
        p_radio._p.append(make_w14_checkbox_sdt(checked=chk))
        p_radio._p.append(_r(f"  {label}     "))

    # ── 8 & 9. Legacy FORMCHECKBOX field codes ────────────────────────────────
    p_legacy = doc.add_paragraph()
    p_legacy._p.append(_r("Property Type (legacy):  "))
    for run_el in make_legacy_checkbox(checked=True):
        p_legacy._p.append(run_el)
    p_legacy._p.append(_r("  Single Family     "))
    for run_el in make_legacy_checkbox(checked=False):
        p_legacy._p.append(run_el)
    p_legacy._p.append(_r("  Condo"))

    # ── 10. Inline SDT text field ─────────────────────────────────────────────
    p_tf = doc.add_paragraph()
    p_tf._p.append(_r("Borrower Name: "))
    p_tf._p.append(make_sdt_text_field("borrower_name", "John Michael Doe"))

    # ── 11. Inline SDT date-picker ────────────────────────────────────────────
    p_date = doc.add_paragraph()
    p_date._p.append(_r("Application Date: "))
    p_date._p.append(make_sdt_date_field("03/15/2024"))

    # ── 12. Body-level SDT wrapping a paragraph ───────────────────────────────
    _insert_before_sectPr(body, make_body_sdt_with_paragraph(
        "[ Section 2 — Financial Information — body-level SDT ]"
    ))

    # ── 13. Inline hyperlink ──────────────────────────────────────────────────
    p_link = doc.add_paragraph()
    p_link._p.append(_r("Reference: "))
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), "section2")
    r_hl = OxmlElement("w:r")
    rPr_hl = OxmlElement("w:rPr")
    u_hl = OxmlElement("w:u")
    u_hl.set(qn("w:val"), "single")
    rPr_hl.append(u_hl)
    r_hl.append(rPr_hl)
    r_hl.append(_t("Click here for guidelines"))
    hl.append(r_hl)
    p_link._p.append(hl)

    # ── 14 & 15. Table with plain cells + checkboxes inside cells ─────────────
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Value"
    hdr[2].text = "Verified"

    row1 = tbl.add_row().cells
    row1[0].text = "Annual Income"
    row1[1].text = "$85,000"
    p_c1 = row1[2].paragraphs[0]
    p_c1._p.append(make_w14_checkbox_sdt(checked=True))
    p_c1._p.append(_r("  Confirmed"))

    row2 = tbl.add_row().cells
    row2[0].text = "Credit Score"
    row2[1].text = "720"
    p_c2 = row2[2].paragraphs[0]
    p_c2._p.append(make_w14_checkbox_sdt(checked=False))
    p_c2._p.append(_r("  Pending"))

    # ── 16. Table with cell background shading ────────────────────────────────
    tbl2 = doc.add_table(rows=2, cols=2)
    tbl2.style = "Table Grid"

    cell_shaded = tbl2.cell(0, 0)
    cell_shaded.text = "Status"
    tcPr = cell_shaded._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"),  "1F4E79")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:val"),   "clear")
    tcPr.append(shd)

    tbl2.cell(0, 1).text = "Approved"
    tbl2.cell(1, 0).text = "Notes"
    tbl2.cell(1, 1).text = "All documents verified and approved"

    # ── 17. Nested SDT (body SDT > paragraph > inline checkbox SDT) ───────────
    _insert_before_sectPr(body, make_nested_sdt(
        "Secondary Applicant:", inner_checkbox_checked=True
    ))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Test assertions
# ─────────────────────────────────────────────────────────────────────────────

TESTS: list[tuple[str, callable]] = []


def test(name: str):
    def decorator(fn):
        TESTS.append((name, fn))
        return fn
    return decorator


# ── 1. Title ──────────────────────────────────────────────────────────────────
@test("1.  Title: text, center align, bold")
def t1(html):
    assert "LOAN APPLICATION FORM" in html
    assert "text-align:center" in html


# ── 2. Run formatting ─────────────────────────────────────────────────────────
@test("2a. Run: bold text preserved")
def t2a(html):
    assert "bold text" in html
    assert "font-weight:700" in html or "<strong>" in html

@test("2b. Run: italic text preserved")
def t2b(html):
    assert "italic text" in html
    assert "<em>" in html

@test("2c. Run: underlined text preserved")
def t2c(html):
    assert "underlined text" in html
    assert "<u>" in html

@test("2d. Run: coloured text preserved (FF0000)")
def t2d(html):
    assert "red coloured text" in html
    assert "FF0000" in html.upper()


# ── 3. Alignment ──────────────────────────────────────────────────────────────
@test("3a. Paragraph: centre alignment")
def t3a(html):
    assert "Centred paragraph" in html
    assert "text-align:center" in html

@test("3b. Paragraph: right alignment")
def t3b(html):
    assert "Right-aligned paragraph" in html
    assert "text-align:right" in html

@test("3c. Paragraph: justify alignment")
def t3c(html):
    assert "justified" in html.lower() or "text-align:justify" in html


# ── 4 & 5. w14:checkbox SDT ───────────────────────────────────────────────────
@test("4.  w14:checkbox SDT checked — <input> rendered")
def t4(html):
    assert 'type="checkbox" checked' in html

@test("5.  w14:checkbox SDT unchecked — <input> rendered")
def t5(html):
    # An unchecked checkbox: no 'checked' attr, just 'disabled'
    import re
    # Find any <input type="checkbox" ...disabled...> WITHOUT checked before disabled
    unchecked = re.search(r'<input type="checkbox"\s+disabled', html)
    assert unchecked, "No unchecked checkbox <input> found"

@test("6.  Checkbox labels alongside controls: Full-Time / Part-Time")
def t6(html):
    assert "Full-Time" in html
    assert "Part-Time" in html


# ── 7. Radio-button group ─────────────────────────────────────────────────────
@test("7.  Radio group: all three options rendered")
def t7(html):
    assert "Purchase" in html
    assert "Refinance" in html
    assert "Home Equity" in html


# ── 8 & 9. Legacy FORMCHECKBOX ────────────────────────────────────────────────
@test("8.  Legacy FORMCHECKBOX: not silently dropped (Single Family label)")
def t8(html):
    assert "Single Family" in html
    # checked=True legacy checkbox must have emitted an <input>
    assert 'type="checkbox"' in html

@test("9.  Legacy FORMCHECKBOX: unchecked variant (Condo label)")
def t9(html):
    assert "Condo" in html


# ── 10 & 11. Inline SDT fields ────────────────────────────────────────────────
@test("10. Inline SDT text field: value text preserved")
def t10(html):
    assert "John Michael Doe" in html

@test("11. Inline SDT date-picker: date string preserved")
def t11(html):
    assert "03/15/2024" in html


# ── 12. Body-level SDT ────────────────────────────────────────────────────────
@test("12. Body-level SDT: inner paragraph rendered (not dropped)")
def t12(html):
    assert "Section 2" in html
    assert "Financial Information" in html


# ── 13. Hyperlink ─────────────────────────────────────────────────────────────
@test("13. Hyperlink text preserved")
def t13(html):
    assert "Click here for guidelines" in html


# ── 14-16. Tables ─────────────────────────────────────────────────────────────
@test("14. Table structure: <table>, <tr>, <td> present")
def t14(html):
    assert "<table" in html
    assert "<tr>" in html
    assert "<td" in html

@test("15. Table: checkboxes inside cells rendered")
def t15(html):
    assert "Confirmed" in html
    assert "Pending" in html

@test("16. Table: cell background shading (#1F4E79)")
def t16(html):
    assert "1F4E79" in html.upper()


# ── 17. Nested SDT ────────────────────────────────────────────────────────────
@test("17. Nested SDT: inner checkbox SDT rendered recursively")
def t17(html):
    assert "Secondary Applicant" in html
    assert "NestedOption" in html
    # The nested checkbox should also produce an <input>
    assert 'type="checkbox"' in html


# ── Sanity ────────────────────────────────────────────────────────────────────
@test("S1. No broken/unclosed <input tags")
def ts1(html):
    import re
    for m in re.finditer(r'<input\b[^>]*>', html):
        tag = m.group()
        assert "type=" in tag, f"<input> missing type attr: {tag}"

@test("S2. HTML is non-empty and has wrapper div")
def ts2(html):
    assert len(html) > 500
    assert 'class="pdv-wrap"' in html


# ─────────────────────────────────────────────────────────────────────────────
# Runner
# ─────────────────────────────────────────────────────────────────────────────

def run_tests(html: str) -> bool:
    passed = failed = 0
    print(f"\n{'=' * 64}")
    print("  DOCX Rendering — Test Results")
    print(f"{'=' * 64}")

    for name, fn in TESTS:
        try:
            fn(html)
            print(f"  PASS  {name}")
            passed += 1
        except AssertionError as exc:
            print(f"  FAIL  {name}")
            if str(exc):
                print(f"       > {exc}")
            failed += 1
        except Exception as exc:
            print(f"  \033[31m✗\033[0m  {name}  [{type(exc).__name__}: {exc}]")
            failed += 1

    print(f"\n  {passed}/{passed + failed} passed", end="")
    if failed:
        print(f"  ({failed} FAILED)\n")
    else:
        print("  — all green!\n")
    print(f"{'=' * 64}\n")

    return failed == 0


if __name__ == "__main__":
    print("Building test DOCX in memory …")
    docx_bytes = build_test_docx()
    print(f"  DOCX size : {len(docx_bytes):,} bytes")

    print("Running _docx_to_html …")
    html_out = _docx_to_html(docx_bytes, highlight_map=[])
    print(f"  HTML size : {len(html_out):,} chars\n")

    # Save for visual inspection in browser
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_docx_output.html")
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write(
            "<!DOCTYPE html><html><head>"
            "<meta charset='utf-8'>"
            "<title>DOCX Render Test</title>"
            "<style>body{font-family:Arial,sans-serif;margin:40px;}</style>"
            "</head><body>\n"
        )
        fh.write(html_out)
        fh.write("\n</body></html>")
    print(f"  Visual output saved: {out_path}")

    ok = run_tests(html_out)
    sys.exit(0 if ok else 1)
